"""Export the knowledge layer to a workbook.

The point of this file is that a fact is only useful to an analyst once it is in
a cell, and a fact in a cell is only trustworthy if you can see where it came
from without leaving the sheet. So every value cell carries its evidence quote,
document and page as a cell comment, and the reconciliation tab explains each
relationship in the same row as the two figures it relates.
"""

import io
import json

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import db, guardrails, security

BODY = Font(name="Arial", size=10)
HEAD = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="27313D")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(bottom=THIN)

TYPE_FILL = {
    "corroborates": PatternFill("solid", fgColor="E4F0E8"),
    "contradicts": PatternFill("solid", fgColor="FBE3E0"),
    "reconciled": PatternFill("solid", fgColor="FBF1DC"),
    "related": PatternFill("solid", fgColor="E7ECF3"),
}


def _safe(row):
    """Untrusted document text is heading for a cell, so neutralise formulas."""
    return [security.safe_cell(guardrails.mask_pii(v)) for v in row]


def _header(ws, headers, widths):
    ws.append(headers)
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=1, column=i)
        c.font, c.fill = HEAD, HEAD_FILL
        c.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def _style_row(ws, row, ncols, fill=None):
    for i in range(1, ncols + 1):
        c = ws.cell(row=row, column=i)
        c.font = BODY
        c.border = BORDER
        c.alignment = Alignment(vertical="top", wrap_text=(i > 2))
        if fill:
            c.fill = fill


def _evidence_note(quote, doc, page, extra=""):
    text = f"{doc} · page {page}\n\n{guardrails.mask_pii(quote)}"
    if extra:
        text = f"{extra}\n\n{text}"
    note = Comment(text[:2000], "Fact knowledge layer")
    note.width, note.height = 420, 200
    return note


def build() -> bytes:
    wb = Workbook()

    # ---------------------------------------------------------------- Read me
    ws = wb.active
    ws.title = "Read me"
    ws.column_dimensions["A"].width = 100
    lines = [
        "Fact knowledge layer export",
        "",
        "Facts tab: every fact extracted from the uploaded PDFs, with the coordinates",
        "that make it comparable. Hover any value cell to see the exact span it came",
        "from, along with its document and page.",
        "",
        "Reconciliation tab: every pair of facts the system linked, and why. Green is",
        "corroboration, red is a contradiction with no context to explain it, amber is",
        "an apparent conflict resolved by a period, basis or scope difference.",
        "",
        "Rejected tab: facts the verifier threw out because the value did not appear in",
        "its own evidence span. This is the failure log, kept deliberately.",
        "",
        "Nothing here is typed by hand. Every value traces back to a page in a source PDF.",
    ]
    for i, line in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=line)
        c.font = Font(name="Arial", size=10, bold=(i == 1))
        c.alignment = Alignment(vertical="top")

    # ----------------------------------------------------------------- Facts
    ws = wb.create_sheet("Facts")
    heads = ["Entity", "Metric", "Value", "Unit", "Period", "Basis", "Scope",
             "Document", "Page"]
    _header(ws, heads, [26, 40, 14, 16, 18, 14, 18, 34, 7])

    facts = db.rows(
        """SELECT f.*, d.filename FROM facts f JOIN documents d ON d.id = f.doc_id
           WHERE f.verified = 1
           ORDER BY f.entity_key, f.metric_key, f.period_start"""
    )
    for f in facts:
        value = f["state_value"] if f["fact_kind"] == "state" else f["value_text"]
        ws.append(_safe([
            f["entity"], f["metric"], value, f["unit"], f["period"],
            f["basis"], f["scope"], f["filename"], f["page"],
        ]))
        r = ws.max_row
        _style_row(ws, r, len(heads))
        qualifiers = json.loads(f["qualifiers"] or "{}")
        extra = ("Qualifiers: " + ", ".join(f"{k}: {v}" for k, v in qualifiers.items())
                 if qualifiers else "")
        ws.cell(row=r, column=3).comment = _evidence_note(
            f["evidence_quote"], f["filename"], f["page"], extra)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(heads))}{max(ws.max_row, 1)}"

    # -------------------------------------------------------- Reconciliation
    ws = wb.create_sheet("Reconciliation")
    heads = ["Relationship", "Entity", "Metric", "Figure A", "Context A",
             "Figure B", "Context B", "Axis that differs", "Why"]
    _header(ws, heads, [16, 22, 32, 16, 26, 16, 26, 18, 70])

    rels = db.rows(
        """SELECT r.*, a.entity a_entity, a.metric a_metric, a.value_text a_value,
                  a.state_value a_state, a.unit a_unit, a.period a_period,
                  a.basis a_basis, a.scope a_scope, a.page a_page,
                  a.evidence_quote a_quote, da.filename a_doc,
                  b.value_text b_value, b.state_value b_state, b.unit b_unit,
                  b.period b_period, b.basis b_basis, b.scope b_scope,
                  b.page b_page, b.evidence_quote b_quote, db_.filename b_doc
           FROM relations r
           JOIN facts a ON a.id = r.fact_a
           JOIN facts b ON b.id = r.fact_b
           JOIN documents da ON da.id = a.doc_id
           JOIN documents db_ ON db_.id = b.doc_id
           ORDER BY CASE r.rel_type WHEN 'contradicts' THEN 0 WHEN 'reconciled' THEN 1
                                    WHEN 'corroborates' THEN 2 ELSE 3 END,
                    r.cross_doc DESC"""
    )
    for r in rels:
        a_val = r["a_state"] or f"{r['a_value']} {r['a_unit'] or ''}".strip()
        b_val = r["b_state"] or f"{r['b_value']} {r['b_unit'] or ''}".strip()
        ctx = lambda s: " · ".join(
            x for x in [r[s + "_period"], r[s + "_basis"], r[s + "_scope"]] if x)
        ws.append(_safe([
            r["rel_type"], r["a_entity"], r["a_metric"],
            a_val, ctx("a"), b_val, ctx("b"),
            ", ".join(json.loads(r["axes_differ"] or "[]")) or "none",
            r["explanation"],
        ]))
        row = ws.max_row
        _style_row(ws, row, len(heads), TYPE_FILL.get(r["rel_type"]))
        ws.cell(row=row, column=4).comment = _evidence_note(
            r["a_quote"], r["a_doc"], r["a_page"])
        ws.cell(row=row, column=6).comment = _evidence_note(
            r["b_quote"], r["b_doc"], r["b_page"])
        if r["support"]:
            s = json.loads(r["support"])
            ws.cell(row=row, column=9).comment = _evidence_note(
                s["evidence_quote"], s["metric"], s["page"], s["statement"])
    ws.auto_filter.ref = f"A1:{get_column_letter(len(heads))}{max(ws.max_row, 1)}"

    # -------------------------------------------------------------- Rejected
    ws = wb.create_sheet("Rejected")
    heads = ["Why it was rejected", "Claimed metric", "Claimed value",
             "Document", "Page", "Span it claimed to come from"]
    _header(ws, heads, [36, 34, 16, 32, 7, 70])
    for f in db.rows(
        """SELECT f.metric, f.value_text, f.unit, f.page, f.evidence_quote,
                  f.reject_reason, d.filename
           FROM facts f JOIN documents d ON d.id = f.doc_id
           WHERE f.verified = 0 ORDER BY f.reject_reason"""
    ):
        ws.append(_safe([f["reject_reason"], f["metric"],
                         f"{f['value_text']} {f['unit'] or ''}".strip(),
                         f["filename"], f["page"], f["evidence_quote"]]))
        _style_row(ws, ws.max_row, len(heads))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
